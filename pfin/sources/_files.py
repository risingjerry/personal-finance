"""Readers for the file shapes platforms export: CSV (UTF-8 or GB18030, with preambles), XLSX, ZIP of CSVs."""
import csv,io
from ..store import zip_members


def csv_table(data,header_at=None):
 """(header, rows with their 1-based line numbers). header_at(rows) -> index of the header row.
 Duplicate headers stay distinct (时间, 时间 [2]); trailing empty cells are dropped."""
 try:text=data.decode('utf-8-sig')
 except UnicodeDecodeError:text=data.decode('gb18030')
 rows=list(csv.reader(io.StringIO(text)));offset=header_at(rows) if header_at else 0
 if len(rows)<=offset:return [],[]
 seen={};header=[]
 for h in rows[offset]:
  h=h.lstrip('﻿').strip();seen[h]=seen.get(h,0)+1;header.append(h if seen[h]==1 else f'{h} [{seen[h]}]')
 out=[]
 for index,row in enumerate(rows[offset+1:],offset+2):
  if not row or not any(row) or row==['没有符合条件的数据。']:continue
  if len(row)>len(header) and not any(row[len(header):]):row=row[:len(header)]
  if len(row)!=len(header):raise ValueError('CSV shape mismatch')
  out.append((index,dict(zip(header,row))))
 return header,out


def csv_files(path,data):
 """(member name, bytes) for a .csv or every CSV inside a .zip/.download."""
 if path.suffix in ('.zip','.download'):return zip_members(data)
 if path.suffix=='.csv':return [(path.name,data)]
 return []


def xlsx_rows(data,header_has):
 """(sheet title, line, row dict) after the first row that contains header_has. Needs openpyxl."""
 from openpyxl import load_workbook
 w=load_workbook(io.BytesIO(data),read_only=True,data_only=True);header=None
 try:
  for index,row in enumerate(w.active.iter_rows(values_only=True),1):
   if header is None:
    if header_has in row:header=list(row)
    continue
   if not row or not row[0]:continue
   yield w.active.title,index,dict(zip(header,row))
 finally:w.close()
